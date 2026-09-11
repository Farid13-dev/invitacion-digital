#!/usr/bin/env python3
"""Genera y envía las invitaciones personalizadas por WhatsApp.

Lee un listado de invitados (xlsx o csv), construye para cada uno un enlace
único y firmado, y abre WhatsApp Web con el mensaje ya redactado.

    python enviar_invitaciones.py --archivo data/invitados.xlsx

El listado necesita dos columnas, `nombre` y `telefono`:

    nombre                                    | telefono
    ------------------------------------------|------------
    Camila Ospina                             | 3001234567
    Andres Betancur & Laura Mejia (Tomas)     | 3007654321

El formato del nombre no es decorativo: "&" separa titulares y los paréntesis
marcan acompañantes. De ahí salen las casillas individuales del formulario de
confirmación, así que un grupo puede responder "vamos dos de los tres".

Esa sintaxis no llega al mensaje de WhatsApp: el saludo se construye aparte
("¡Hola Andres y Laura!"), y el grupo completo queda disponible como {grupo}.

Antes de mandar nada de verdad:

    python enviar_invitaciones.py --archivo data/invitados.xlsx --modo enlaces

genera `enlaces-generados.csv` con todos los enlaces y no abre ni una pestaña.

Ver README.md en esta carpeta para el resto.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import hmac
import os
import re
import sys
import time
import unicodedata
import urllib.parse
import webbrowser
from dataclasses import dataclass
from pathlib import Path

# ─────────────────────────────────────────────────────────────────────────
# Constantes
# ─────────────────────────────────────────────────────────────────────────

LONGITUD_FIRMA = 12  # debe coincidir con CONFIG.LONGITUD_FIRMA del Apps Script
# Marcadores que acepta la plantilla del mensaje, con valores de mentira para
# validarla antes de empezar a enviar.
MARCADORES_DE_PRUEBA = {"nombre": "", "grupo": "", "enlace": ""}
# Titulares que no son una persona: de estos se saluda el nombre completo
# ("Familia Betancur"), no la primera palabra.
COLECTIVOS = {"familia", "flia", "fam", "sres", "señores", "senores", "hogar"}
COLUMNAS = ("nombre", "telefono")
SALIDA_ENLACES = "enlaces-generados.csv"

# WhatsApp Web tarda en cargar la conversación. Menos de ~12 s y se corre el
# riesgo de que la siguiente pestaña se abra antes de haber enviado la anterior.
ESPERA_POR_DEFECTO = 15

PLANTILLA_POR_DEFECTO = (
    "¡Hola {nombre}! 🤍\n\n"
    "Con mucha alegría queremos compartir contigo uno de los momentos más "
    "importantes de nuestras vidas. Hemos preparado una invitación especial para ti:\n\n"
    "{enlace}\n\n"
    "Tu presencia sería un regalo muy especial para nosotros.\n\n"
    "¡Esperamos celebrar juntos este día inolvidable! ✨"
)


# ─────────────────────────────────────────────────────────────────────────
# Modelo
# ─────────────────────────────────────────────────────────────────────────


@dataclass
class Invitado:
    fila: int
    nombre: str
    telefono: str

    @property
    def valido(self) -> bool:
        return bool(self.nombre) and 7 <= len(self.telefono) <= 15

    @property
    def saludo(self) -> str:
        """Cómo se le habla a este grupo en el mensaje de WhatsApp.

        `nombre` lleva la sintaxis del grupo ("Andrés Betancur & Laura Mejía
        (Tomás, Sara)"), que es justo lo que la invitación necesita para
        generar las casillas — pero un saludo de WhatsApp no puede decir
        "¡Hola Andrés Betancur & Laura Mejía (Tomás, Sara)!".

        Aquí se queda con los titulares y con el nombre de pila:
        "¡Hola Andrés y Laura!".
        """
        titulares = [nombre_de_trato(t) for t in separar_titulares(self.nombre)]
        if not titulares:
            return self.nombre
        if len(titulares) == 1:
            return titulares[0]
        return ", ".join(titulares[:-1]) + " y " + titulares[-1]


# ─────────────────────────────────────────────────────────────────────────
# Lectura del listado
# ─────────────────────────────────────────────────────────────────────────


def separar_titulares(nombre: str) -> list[str]:
    """Los titulares del grupo, sin los acompañantes entre paréntesis.

    Es el mismo formato que descompone `src/lib/invitados.ts` en el navegador;
    aquí solo hace falta la parte de los titulares, para el saludo.
    """
    base = re.sub(r"\([^)]*\)", "", nombre)
    return [p.strip() for p in re.split(r"[&,]", base) if p.strip()]


def nombre_de_trato(titular: str) -> str:
    """El nombre de pila, salvo cuando el titular no es una persona.

    "Camila Ospina" se saluda como "Camila", pero "Familia Betancur" no se
    saluda como "Familia": ahí el apellido es justo lo que hace falta.
    """
    palabras = titular.split()
    if not palabras:
        return titular
    if palabras[0].rstrip(".").lower() in COLECTIVOS:
        return titular
    return palabras[0]


def limpiar_telefono(valor: object) -> str:
    """Deja solo dígitos.

    Excel guarda los teléfonos con marcas de dirección invisibles (U+202A,
    U+202C) que se cuelan al copiar desde WhatsApp. A la vista no existen,
    pero hacen que el número no coincida con el de la hoja de cálculo.
    """
    if valor is None:
        return ""
    texto = unicodedata.normalize("NFKC", str(valor))
    return re.sub(r"\D", "", texto)


def limpiar_nombre(valor: object) -> str:
    if valor is None:
        return ""
    # Los espacios dobles y los finales vienen casi siempre del copiar/pegar,
    # y acabarían mostrándose tal cual en la invitación.
    return re.sub(r"\s+", " ", str(valor)).strip()


def leer_listado(ruta: Path) -> list[Invitado]:
    if not ruta.exists():
        raise SystemExit(f"No existe el archivo: {ruta}")

    if ruta.suffix.lower() == ".csv":
        filas = _leer_csv(ruta)
    elif ruta.suffix.lower() in (".xlsx", ".xlsm"):
        filas = _leer_xlsx(ruta)
    else:
        raise SystemExit(f"Formato no soportado: {ruta.suffix}. Usa .xlsx o .csv")

    invitados = [
        Invitado(
            fila=numero,
            nombre=limpiar_nombre(fila.get("nombre")),
            telefono=limpiar_telefono(fila.get("telefono")),
        )
        for numero, fila in filas
    ]

    _avisar_de_duplicados(invitados)
    return invitados


def _leer_csv(ruta: Path) -> list[tuple[int, dict]]:
    # utf-8-sig se come el BOM que Excel escribe al exportar a CSV y que si no
    # convierte la primera columna en "﻿nombre".
    with ruta.open(encoding="utf-8-sig", newline="") as f:
        lector = csv.DictReader(f)
        _validar_columnas(lector.fieldnames or [], ruta)
        return [(i, fila) for i, fila in enumerate(lector, start=2)]


def _leer_xlsx(ruta: Path) -> list[tuple[int, dict]]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise SystemExit(
            "Para leer archivos .xlsx hace falta openpyxl:\n"
            "    pip install -r requirements.txt\n"
            "(o exporta el listado a .csv desde Excel)"
        )

    libro = load_workbook(ruta, read_only=True, data_only=True)

    try:
        hoja = libro.active
        if hoja is None:
            raise SystemExit(f"El archivo {ruta} no tiene ninguna hoja.")

        filas = hoja.iter_rows(values_only=True)

        try:
            encabezados = [limpiar_nombre(c).lower() for c in next(filas)]
        except StopIteration:
            raise SystemExit(f"El archivo {ruta} está vacío.")

        _validar_columnas(encabezados, ruta)

        return [
            (numero, dict(zip(encabezados, valores)))
            for numero, valores in enumerate(filas, start=2)
        ]
    finally:
        # read_only deja un descriptor abierto; sin esto, un error a mitad de
        # la lectura lo mantiene tomado hasta que muere el proceso.
        libro.close()


def _validar_columnas(encabezados: list[str], ruta: Path) -> None:
    presentes = {str(h).strip().lower() for h in encabezados if h}
    faltan = [c for c in COLUMNAS if c not in presentes]
    if faltan:
        raise SystemExit(
            f"A {ruta.name} le faltan columnas: {', '.join(faltan)}.\n"
            f"Encontradas: {', '.join(sorted(presentes)) or '(ninguna)'}"
        )


def _avisar_de_duplicados(invitados: list[Invitado]) -> None:
    """Un teléfono repetido significa que alguien recibirá dos invitaciones y
    que la segunda confirmación pisará a la primera en la hoja."""
    vistos: dict[str, int] = {}
    for inv in invitados:
        if not inv.telefono:
            continue
        if inv.telefono in vistos:
            print(
                f"  [!] Telefono repetido {inv.telefono}: "
                f"filas {vistos[inv.telefono]} y {inv.fila}",
                file=sys.stderr,
            )
        else:
            vistos[inv.telefono] = inv.fila


# ─────────────────────────────────────────────────────────────────────────
# Enlaces
# ─────────────────────────────────────────────────────────────────────────


def firmar(telefono: str, secreto: str) -> str:
    """HMAC-SHA256 del teléfono, truncado.

    Idéntico a `firmaEsperada()` del Apps Script. Es lo que impide que alguien
    lea o sobrescriba la confirmación de otro invitado probando teléfonos: el
    secreto no sale nunca de aquí ni de las propiedades del script.
    """
    digest = hmac.new(secreto.encode("utf-8"), telefono.encode("utf-8"), hashlib.sha256)
    return digest.hexdigest()[:LONGITUD_FIRMA]


def construir_enlace(base: str, invitado: Invitado, secreto: str) -> str:
    """Enlace único del invitado: ?inv=<nombres>&tel=<telefono>&f=<firma>.

    Se construye con urlsplit/urlunsplit en vez de concatenando cadenas para
    que siga funcionando si la URL base ya trae su propia query o no termina
    en barra.
    """
    partes = urllib.parse.urlsplit(base)

    consulta = urllib.parse.parse_qsl(partes.query)
    consulta.append(("inv", invitado.nombre))
    consulta.append(("tel", invitado.telefono))
    if secreto:
        consulta.append(("f", firmar(invitado.telefono, secreto)))

    return urllib.parse.urlunsplit(
        (partes.scheme, partes.netloc, partes.path or "/", urllib.parse.urlencode(consulta), "")
    )


def construir_url_whatsapp(telefono: str, indicativo: str, mensaje: str) -> str:
    # Un móvil colombiano tiene 10 dígitos. Si el listado trae más, ya incluía
    # el indicativo y volver a anteponerlo dejaría el número inservible.
    numero = telefono if (not indicativo or len(telefono) > 10) else f"{indicativo}{telefono}"

    return "https://web.whatsapp.com/send" f"?phone={numero}&text={urllib.parse.quote(mensaje)}"


def cargar_plantilla(ruta: Path | None) -> str:
    if ruta is None:
        return PLANTILLA_POR_DEFECTO
    if not ruta.exists():
        raise SystemExit(f"No existe la plantilla de mensaje: {ruta}")

    plantilla = ruta.read_text(encoding="utf-8")

    if "{enlace}" not in plantilla:
        raise SystemExit(
            f"{ruta} no contiene {{enlace}}: el invitado recibiría un mensaje "
            "sin su invitación."
        )

    # Una llave suelta en el texto ("{" de un emoji copiado, una nota entre
    # llaves) hace que .format() reviente a mitad del envío, con cuarenta
    # pestañas ya abiertas. Se comprueba aquí, antes de mandar nada.
    try:
        plantilla.format(**MARCADORES_DE_PRUEBA)
    except (KeyError, IndexError, ValueError) as error:
        raise SystemExit(
            f"{ruta} tiene un marcador que no se puede rellenar: {error}\n"
            f"Los válidos son {{nombre}}, {{grupo}} y {{enlace}}. Para escribir "
            "una llave literal, dóblala: {{{{ }}}}"
        )

    return plantilla


# ─────────────────────────────────────────────────────────────────────────
# Ejecución
# ─────────────────────────────────────────────────────────────────────────


def construir_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(
        prog="enviar_invitaciones.py",
        description="Genera y envía las invitaciones personalizadas por WhatsApp.",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=(
            "Ejemplos:\n"
            "  # Ver los enlaces sin abrir nada (haz esto primero)\n"
            "  python enviar_invitaciones.py -a data/invitados.xlsx --modo enlaces\n\n"
            "  # Enviar de verdad\n"
            "  python enviar_invitaciones.py -a data/invitados.xlsx\n\n"
            "  # Reanudar tras una interrupción en la fila 34\n"
            "  python enviar_invitaciones.py -a data/invitados.xlsx --desde 34\n\n"
            "  # Reenviar a una sola persona\n"
            "  python enviar_invitaciones.py -a data/invitados.xlsx --solo 3001234567\n"
        ),
    )

    p.add_argument(
        "-a", "--archivo", type=Path, required=True,
        help="listado de invitados (.xlsx o .csv) con columnas nombre y telefono",
    )
    p.add_argument(
        "-u", "--url", default=os.environ.get("INVITACION_URL"),
        help="URL base de la invitación (por defecto: INVITACION_URL del entorno)",
    )
    p.add_argument(
        "--modo", choices=("whatsapp", "enlaces", "prueba"), default="whatsapp",
        help=(
            "whatsapp: abre WhatsApp Web con el mensaje listo (por defecto) | "
            "enlaces: solo escribe el CSV de enlaces | "
            "prueba: abre la invitación en el navegador, sin WhatsApp"
        ),
    )
    p.add_argument(
        "-m", "--mensaje", type=Path, default=None,
        help="plantilla de texto con {nombre} y {enlace} (por defecto: mensaje.txt si existe)",
    )
    p.add_argument(
        "-i", "--indicativo", default=os.environ.get("INDICATIVO_PAIS", "57"),
        help="indicativo del país sin '+' (por defecto: 57, Colombia)",
    )
    p.add_argument(
        "-e", "--espera", type=float, default=ESPERA_POR_DEFECTO,
        help=f"segundos entre envíos (por defecto: {ESPERA_POR_DEFECTO})",
    )
    p.add_argument("--desde", type=int, default=0, metavar="N",
                   help="empezar en la fila N del listado (para reanudar)")
    p.add_argument("--limite", type=int, default=None, metavar="N",
                   help="enviar como mucho N invitaciones")
    p.add_argument("--solo", default=None, metavar="TEL",
                   help="enviar únicamente a este teléfono")
    p.add_argument("--sin-firma", action="store_true",
                   help="generar enlaces sin firma HMAC (solo para pruebas locales)")

    return p


def cargar_dotenv() -> None:
    """Carga `.env` del proyecto sin depender de python-dotenv."""
    for candidato in (Path(".env"), Path(__file__).resolve().parent.parent / ".env"):
        if not candidato.exists():
            continue
        for linea in candidato.read_text(encoding="utf-8").splitlines():
            linea = linea.strip()
            if not linea or linea.startswith("#") or "=" not in linea:
                continue
            clave, _, valor = linea.partition("=")
            os.environ.setdefault(clave.strip(), valor.strip().strip('"').strip("'"))
        return


def main(argv: list[str] | None = None) -> int:
    # La consola de Windows usa cp1252 por defecto y revienta al imprimir
    # nombres con caracteres fuera de esa tabla. Preferimos un interrogante
    # en pantalla antes que un UnicodeEncodeError a mitad de un envío.
    for flujo in (sys.stdout, sys.stderr):
        if hasattr(flujo, "reconfigure"):
            flujo.reconfigure(encoding="utf-8", errors="replace")

    cargar_dotenv()
    args = construir_parser().parse_args(argv)

    if not args.url:
        raise SystemExit(
            "Falta la URL de la invitación.\n"
            "Pásala con --url o define INVITACION_URL en el archivo .env"
        )

    secreto = "" if args.sin_firma else os.environ.get("SECRETO_RSVP", "")
    if not secreto and not args.sin_firma:
        raise SystemExit(
            "Falta SECRETO_RSVP en el entorno.\n"
            "Debe ser el mismo valor que la propiedad SECRETO_RSVP del Apps Script.\n"
            "Para probar sin firma: --sin-firma"
        )

    plantilla_ruta = args.mensaje
    if plantilla_ruta is None:
        por_defecto = Path(__file__).resolve().parent / "mensaje.txt"
        plantilla_ruta = por_defecto if por_defecto.exists() else None
    plantilla = cargar_plantilla(plantilla_ruta)

    invitados = leer_listado(args.archivo)

    # ── Filtros ──────────────────────────────────────────────────────────
    validos = [i for i in invitados if i.valido]
    descartados = [i for i in invitados if not i.valido]

    if args.solo:
        buscado = limpiar_telefono(args.solo)
        validos = [i for i in validos if i.telefono == buscado]
        if not validos:
            raise SystemExit(f"No hay ningún invitado con el teléfono {buscado}.")

    if args.desde:
        validos = [i for i in validos if i.fila >= args.desde]

    if args.limite is not None:
        validos = validos[: args.limite]

    # ── Resumen ──────────────────────────────────────────────────────────
    print(f"\n  Listado      {args.archivo}")
    print(f"  Invitación   {args.url}")
    print(f"  Modo         {args.modo}")
    print(f"  Firma        {'sí' if secreto else 'NO (--sin-firma)'}")
    print(f"  A enviar     {len(validos)} de {len(invitados)} filas")

    if descartados:
        print(f"\n  {len(descartados)} fila(s) descartadas por datos incompletos:")
        for inv in descartados[:10]:
            motivo = "sin nombre" if not inv.nombre else f"teléfono inválido ({inv.telefono or '—'})"
            print(f"    fila {inv.fila}: {inv.nombre or '—'} — {motivo}")
        if len(descartados) > 10:
            print(f"    … y {len(descartados) - 10} más")

    if not validos:
        print("\n  Nada que enviar.\n")
        return 0

    # ── Enlaces (siempre se escriben, sirven de respaldo para reenviar) ──
    enlaces = [(inv, construir_enlace(args.url, inv, secreto)) for inv in validos]
    ruta_csv = Path(SALIDA_ENLACES)

    with ruta_csv.open("w", encoding="utf-8-sig", newline="") as f:
        escritor = csv.writer(f)
        escritor.writerow(["fila", "nombre", "telefono", "enlace"])
        for inv, enlace in enlaces:
            escritor.writerow([inv.fila, inv.nombre, inv.telefono, enlace])

    print(f"\n  Enlaces escritos en {ruta_csv}")

    if args.modo == "enlaces":
        print("  Modo 'enlaces': no se abrió ninguna pestaña.\n")
        return 0

    # ── Confirmación antes de abrir 80 pestañas ─────────────────────────
    if args.modo == "whatsapp":
        minutos = len(enlaces) * args.espera / 60
        print(f"  Duración estimada: ~{minutos:.0f} min\n")
        respuesta = input("  ¿Abrir WhatsApp Web y empezar? [s/N] ").strip().lower()
        if respuesta not in ("s", "si", "sí", "y"):
            print("  Cancelado.\n")
            return 0

    print()
    return enviar(enlaces, args, plantilla)


def enviar(
    enlaces: list[tuple[Invitado, str]],
    args: argparse.Namespace,
    plantilla: str,
) -> int:
    total = len(enlaces)
    ultima_fila = 0

    try:
        for indice, (inv, enlace) in enumerate(enlaces, start=1):
            ultima_fila = inv.fila
            etiqueta = f"[{indice}/{total}] fila {inv.fila}  {inv.nombre}"

            if args.modo == "prueba":
                print(f"  {etiqueta}\n      {enlace}")
                webbrowser.open(enlace)
                time.sleep(1.5)
                continue

            mensaje = plantilla.format(nombre=inv.saludo, grupo=inv.nombre, enlace=enlace)
            webbrowser.open(construir_url_whatsapp(inv.telefono, args.indicativo, mensaje))
            print(f"  {etiqueta}  ({inv.telefono})")

            if indice < total:
                time.sleep(args.espera)

    except KeyboardInterrupt:
        # Se reanuda en la fila interrumpida, no en la siguiente: no hay forma
        # de saber si esa llegó a enviarse. Repetir una invitación es
        # incómodo; olvidarse de un invitado es un asiento vacío el día de la
        # boda que nadie sabe explicar.
        print(
            f"\n\n  Interrumpido en la fila {ultima_fila}.\n"
            f"  Para continuar (comprueba si esa última llegó a enviarse):\n"
            f"      python enviar_invitaciones.py -a {args.archivo} --desde {ultima_fila}\n"
        )
        return 130

    print(f"\n  Listo: {total} invitaciones abiertas.\n")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
